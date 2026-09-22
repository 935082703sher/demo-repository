# -*- coding: utf-8 -*-
"""
Manba hujjatlarni (docx / doc / pdf) matnga aylantiradi va raw/ papkasiga yozadi.
Shundan keyin build_letters.py va build_kb.py ishga tushiriladi.

Foydalanish:
    python3 ingest_sources.py /yo'l/manba_hujjatlar_papkasi

.doc (eski format) uchun `libreoffice` yoki `antiword` kerak:
    sudo apt install libreoffice-writer   # yoki: antiword
"""
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "raw"


def from_docx(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    parts = [par.text for par in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def from_pdf(p: Path) -> str:
    import pdfplumber
    out = []
    with pdfplumber.open(str(p)) as pdf:
        for page in pdf.pages:
            out.append(page.extract_text() or "")
    text = "\n".join(out)
    if len(text.strip()) < 100:
        print(f"  ! {p.name}: matn topilmadi — skan bo'lishi mumkin, OCR kerak "
              f"(tesseract-ocr + uzb/rus til paketi)")
    return text


def from_doc(p: Path) -> str:
    for cmd in (["antiword", str(p)],
                ["libreoffice", "--headless", "--convert-to", "txt:Text",
                 "--outdir", str(RAW), str(p)]):
        try:
            r = subprocess.run(cmd, capture_output=True, timeout=120)
            if cmd[0] == "antiword" and r.returncode == 0:
                return r.stdout.decode("utf-8", "ignore")
            if cmd[0] == "libreoffice" and r.returncode == 0:
                txt = RAW / (p.stem + ".txt")
                if txt.exists():
                    return txt.read_text(encoding="utf-8", errors="ignore")
        except (FileNotFoundError, subprocess.TimeoutExpired):
            continue
    print(f"  ! {p.name}: .doc o'qilmadi — libreoffice yoki antiword o'rnating")
    return ""


HANDLERS = {".docx": from_docx, ".pdf": from_pdf, ".doc": from_doc,
            ".txt": lambda p: p.read_text(encoding="utf-8", errors="ignore")}


def main(src_dir: str):
    RAW.mkdir(parents=True, exist_ok=True)
    files = sorted(Path(src_dir).rglob("*"))
    n = 0
    for f in files:
        h = HANDLERS.get(f.suffix.lower())
        if not h or not f.is_file():
            continue
        try:
            text = h(f)
        except Exception as e:                      # noqa: BLE001
            print(f"  ! {f.name}: {e}")
            continue
        if not text.strip():
            continue
        (RAW / (f.stem + ".txt")).write_text(text, encoding="utf-8")
        print(f"  ✓ {f.name} → raw/{f.stem}.txt ({len(text)} belgi)")
        n += 1
    print(f"\nJami {n} ta hujjat raw/ ga yozildi.")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1])
